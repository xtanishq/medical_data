"""Read supplied workbooks and generate immutable website records.

Main sheet identity is authoritative; import-tab IDs can contain template IDs.
Clinical values, references, notes and statuses are preserved from the source.
"""
import json
import re
import sys
from pathlib import Path
from openpyxl import load_workbook

PARAMETERS = [
    ('heart-rate', 'Cardiovascular'), ('blood-pressure', 'Cardiovascular'),
    ('cvp', 'Cardiovascular'), ('capillary-refill', 'Cardiovascular'),
    ('respiratory-rate', 'Respiratory'), ('spo2', 'Respiratory'),
    ('etco2', 'Respiratory'), ('ventilator', 'Respiratory'),
    ('gcs', 'Neurological'), ('pupils', 'Neurological'), ('rass', 'Neurological'),
    ('urine-output', 'Fluids & renal'), ('fluid-balance', 'Fluids & renal'),
    ('glucose', 'Metabolic'), ('temperature', 'Metabolic'), ('pain', 'Neurological'),
    ('line-integrity', 'Lines & safety'), ('vasoactive', 'Cardiovascular'),
    ('neurovascular', 'Lines & safety'), ('drain-output', 'Lines & safety'),
]
STATUSES = {'normal': 'Normal', 'average': 'Average', 'review': 'Review',
            'critical': 'Critical', 'not assessed': 'Not Assessed'}

def numeric(parameter, value):
    if parameter in {'capillary-refill', 'pupils', 'ventilator', 'pain',
                     'line-integrity', 'vasoactive', 'neurovascular'}:
        return None
    if parameter == 'blood-pressure':
        match = re.search(r'MAP\s+(-?\d+(?:\.\d+)?)', value)
    else:
        match = re.search(r'[-+]?\d+(?:\.\d+)?', value)
    return float(match.group(1) if parameter == 'blood-pressure' else match.group()) if match else None

def read_patient(path):
    if path.suffix == '.numbers':
        from numbers_parser import Document
        from types import SimpleNamespace

        class Sheet:
            def __init__(self, rows):
                self.rows = rows
            def cell(self, row, col):
                return SimpleNamespace(value=self.rows[row-1][col-1])
            def __getitem__(self, address):
                match = re.fullmatch(r'([A-Z])(\d+)', address)
                return self.cell(int(match[2]), ord(match[1])-64)
            def iter_rows(self, min_row, values_only):
                return iter(self.rows[min_row-1:])

        wb = {s.name: Sheet(s.tables[0].rows(values_only=True)) for s in Document(str(path)).sheets}
    else:
        wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb['24-Hour Entry']
    patient_id = str(ws['B2'].value).strip().upper()
    # User approved GU70051 for Priyanshu to resolve the source's duplicate GU70050.
    if path.name == 'priyanshu.numbers':
        patient_id = 'GU70051'
    assert re.fullmatch(r'[A-Z]{2}\d{5}', patient_id), patient_id
    date = ws['E2'].value.strftime('%Y-%m-%d')
    imports = {(str(row[2]), str(row[3])): row
               for row in wb['App Import Format'].iter_rows(min_row=2, values_only=True)}
    records = []
    for row_index in range(8, 32):
        hour = str(ws.cell(row_index, 1).value)
        assert hour == f'{row_index-8:02}:00', hour
        status = STATUSES[str(ws.cell(row_index, 22).value).strip().lower()]
        observations = []
        for col, (parameter, category) in enumerate(PARAMETERS, 2):
            label = str(ws.cell(6, col).value)
            observed = str(ws.cell(row_index, col).value)
            source = imports[(hour, label)]
            assert observed == str(source[5]), (patient_id, hour, label)
            observations.append(dict(parameterId=parameter, label=label, category=category,
                normalRange=str(ws.cell(7, col).value), observedValue=observed,
                unit=source[6], status=STATUSES[str(source[7]).strip().lower()],
                numericValue=numeric(parameter, observed)))
        records.append(dict(hour=hour, overallStatus=status,
            actionNotes=str(ws.cell(row_index, 23).value or ''), observations=observations))
    name = str(ws['B1'].value).strip()
    photo = next((f'/patients/{key}.jpeg' for key in ['gaurav','jatin','priyanshu']
                  if key in name.lower()), None)
    bed = ws['H2'].value
    if isinstance(bed, float) and bed.is_integer():
        bed = int(bed)
    return dict(patientId=patient_id, displayName=name, bed=str(bed),
        unit='ICU', recordDate=date, lastUpdated=f'{date}T23:00:00+05:30',
        photoUrl=photo, overallStatus=records[-1]['overallStatus'], hourlyRecords=records)

patients = [read_patient(Path(path)) for path in sys.argv[1:]]
assert len({p['patientId'] for p in patients}) == len(patients)
output = Path(__file__).resolve().parents[1] / 'src/icu/icu.imported.ts'
output.write_text("// Generated from supplied patient workbooks by scripts/import-patient-data.py.\n"
    "import type { PatientRecord } from './icu.types';\n\n"
    "export const IMPORTED_PATIENTS: PatientRecord[] = "
    + json.dumps(patients, ensure_ascii=False, indent=2) + ';\n')
print(f'Imported {len(patients)} patients; {len(patients)*24*20} observations verified against import tabs.')
